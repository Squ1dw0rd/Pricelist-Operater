"""
Excel Output Generator for Supplier Price List Consolidation Tool
Handles creation of multi-sheet Excel workbooks with directory and hyperlinks.
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.hyperlink import Hyperlink
import logging
from typing import Dict, List, Optional, Any
from pathlib import Path
from datetime import datetime
import re

class ExcelGenerator:
    """Generates consolidated Excel workbooks with multiple supplier sheets."""
    
    def __init__(self, config: Dict[str, Any]):
        """
        Initialize the Excel generator.
        
        Args:
            config: Configuration dictionary
        """
        self.config = config
        self.output_config = config.get('output', {})
        self.directory_sheet_name = self.output_config.get('directory_sheet_name', 'Directory')
        
        # Styling configuration
        self.header_style = self._create_header_style()
        self.hyperlink_style = self._create_hyperlink_style()
        self.border_style = self._create_border_style()
    
    def generate_master_workbook(self, supplier_data: Dict[str, pd.DataFrame], 
                                output_path: str, validation_reports: Dict[str, Any] = None) -> str:
        """
        Generate the master Excel workbook with all supplier data.
        
        Args:
            supplier_data: Dictionary of supplier_name -> DataFrame
            output_path: Path for the output Excel file
            validation_reports: Optional validation reports for each supplier
        
        Returns:
            Path to the generated Excel file
        """
        logging.info(f"Generating master workbook with {len(supplier_data)} suppliers")
        
        # Create workbook
        workbook = openpyxl.Workbook()
        
        # Remove default sheet
        workbook.remove(workbook.active)
        
        # Create directory sheet first
        directory_sheet = workbook.create_sheet(self.directory_sheet_name)
        
        # Create supplier sheets
        supplier_sheets = {}
        for supplier_name, df in supplier_data.items():
            sheet_name = self._sanitize_sheet_name(supplier_name)
            sheet = workbook.create_sheet(sheet_name)
            supplier_sheets[supplier_name] = sheet
            
            # Populate supplier sheet
            self._populate_supplier_sheet(sheet, df, supplier_name, validation_reports)
        
        # Populate directory sheet
        self._populate_directory_sheet(directory_sheet, supplier_data, supplier_sheets, validation_reports)
        
        # Move directory sheet to first position
        workbook.move_sheet(directory_sheet, offset=-len(workbook.worksheets))
        
        # Save workbook
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        
        workbook.save(output_path)
        logging.info(f"Master workbook saved to: {output_path}")
        
        return str(output_path)
    
    def _populate_directory_sheet(self, sheet, supplier_data: Dict[str, pd.DataFrame], 
                                 supplier_sheets: Dict[str, Any], validation_reports: Dict[str, Any] = None):
        """Populate the directory sheet with supplier links and summary."""
        # Title
        sheet['A1'] = "Supplier Price List Directory"
        sheet['A1'].font = Font(size=16, bold=True)
        sheet.merge_cells('A1:F1')
        
        # Metadata
        current_row = 3
        if self.output_config.get('include_processing_metadata', True):
            sheet[f'A{current_row}'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            sheet[f'A{current_row}'].font = Font(italic=True)
            current_row += 1
            
            sheet[f'A{current_row}'] = f"Total Suppliers: {len(supplier_data)}"
            current_row += 1
            
            total_rows = sum(len(df) for df in supplier_data.values())
            sheet[f'A{current_row}'] = f"Total Products: {total_rows}"
            current_row += 2
        
        # Headers
        headers = ['Supplier Name', 'Products', 'Status', 'Errors', 'Warnings', 'Sheet Link']
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=current_row, column=col, value=header)
            cell.font = self.header_style['font']
            cell.fill = self.header_style['fill']
            cell.alignment = self.header_style['alignment']
            cell.border = self.border_style
        
        current_row += 1
        
        # Supplier data
        for supplier_name, df in supplier_data.items():
            sheet_name = self._sanitize_sheet_name(supplier_name)
            
            # Supplier name
            sheet.cell(row=current_row, column=1, value=supplier_name)
            
            # Product count
            sheet.cell(row=current_row, column=2, value=len(df))
            
            # Status, errors, warnings from validation
            status = "OK"
            error_count = 0
            warning_count = 0
            
            if validation_reports and supplier_name in validation_reports:
                report = validation_reports[supplier_name]
                summary = report.get_summary()
                error_count = summary['error_count']
                warning_count = summary['warning_count']
                
                if error_count > 0:
                    status = "ERRORS"
                elif warning_count > 0:
                    status = "WARNINGS"
            
            status_cell = sheet.cell(row=current_row, column=3, value=status)
            if status == "ERRORS":
                status_cell.font = Font(color="FF0000", bold=True)
            elif status == "WARNINGS":
                status_cell.font = Font(color="FF8C00", bold=True)
            else:
                status_cell.font = Font(color="008000", bold=True)
            
            sheet.cell(row=current_row, column=4, value=error_count)
            sheet.cell(row=current_row, column=5, value=warning_count)
            
            # Hyperlink to supplier sheet
            link_cell = sheet.cell(row=current_row, column=6, value=f"Go to {supplier_name}")
            link_cell.hyperlink = f"#{sheet_name}!A1"
            link_cell.font = self.hyperlink_style['font']
            
            # Apply borders
            for col in range(1, 7):
                sheet.cell(row=current_row, column=col).border = self.border_style
            
            current_row += 1
        
        # Auto-adjust column widths
        self._auto_adjust_columns(sheet)
    
    def _populate_supplier_sheet(self, sheet, df: pd.DataFrame, supplier_name: str, 
                               validation_reports: Dict[str, Any] = None):
        """Populate a supplier sheet with data."""
        # Title
        sheet['A1'] = f"{supplier_name} - Price List"
        sheet['A1'].font = Font(size=14, bold=True)
        sheet.merge_cells(f'A1:{self._get_column_letter(len(df.columns))}1')
        
        # Metadata
        current_row = 3
        if self.output_config.get('include_processing_metadata', True):
            sheet[f'A{current_row}'] = f"Processed: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            sheet[f'A{current_row}'].font = Font(italic=True)
            current_row += 1
            
            sheet[f'A{current_row}'] = f"Total Products: {len(df)}"
            current_row += 1
            
            # Validation summary
            if validation_reports and supplier_name in validation_reports:
                report = validation_reports[supplier_name]
                summary = report.get_summary()
                
                sheet[f'A{current_row}'] = f"Validation: {summary['error_count']} errors, {summary['warning_count']} warnings"
                if summary['error_count'] > 0:
                    sheet[f'A{current_row}'].font = Font(color="FF0000")
                elif summary['warning_count'] > 0:
                    sheet[f'A{current_row}'].font = Font(color="FF8C00")
                else:
                    sheet[f'A{current_row}'].font = Font(color="008000")
                current_row += 1
            
            current_row += 1
        
        # Back to directory link
        back_link = sheet[f'A{current_row}']
        back_link.value = "← Back to Directory"
        back_link.hyperlink = f"#{self.directory_sheet_name}!A1"
        back_link.font = self.hyperlink_style['font']
        current_row += 2
        
        # Data headers
        header_row = current_row
        for col, column_name in enumerate(df.columns, 1):
            cell = sheet.cell(row=header_row, column=col, value=column_name.replace('_', ' ').title())
            cell.font = self.header_style['font']
            cell.fill = self.header_style['fill']
            cell.alignment = self.header_style['alignment']
            cell.border = self.border_style
        
        # Data rows
        for row_idx, (_, row) in enumerate(df.iterrows(), header_row + 1):
            for col_idx, value in enumerate(row, 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                
                # Format value
                if pd.isna(value):
                    cell.value = ""
                elif isinstance(value, (int, float)):
                    cell.value = value
                    # Format currency columns
                    if any(price_col in df.columns[col_idx-1].lower() for price_col in ['price', 'cost']):
                        cell.number_format = '"$"#,##0.00'
                else:
                    cell.value = str(value)
                
                cell.border = self.border_style
                
                # Highlight flagged rows if validation data available
                if validation_reports and supplier_name in validation_reports:
                    report = validation_reports[supplier_name]
                    original_row_idx = row_idx - header_row - 1  # Adjust for header offset
                    if original_row_idx in report.flagged_rows:
                        cell.fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        
        # Auto-adjust column widths
        self._auto_adjust_columns(sheet)
        
        # Freeze header row
        sheet.freeze_panes = f'A{header_row + 1}'
    
    def _sanitize_sheet_name(self, name: str) -> str:
        """Sanitize sheet name to comply with Excel requirements."""
        # Remove invalid characters
        sanitized = re.sub(r'[\\/*?:\[\]]', '_', name)
        
        # Limit length to 31 characters
        if len(sanitized) > 31:
            sanitized = sanitized[:28] + "..."
        
        # Ensure it's not empty
        if not sanitized:
            sanitized = "Sheet"
        
        return sanitized
    
    def _get_column_letter(self, col_num: int) -> str:
        """Convert column number to Excel column letter."""
        result = ""
        while col_num > 0:
            col_num -= 1
            result = chr(col_num % 26 + ord('A')) + result
            col_num //= 26
        return result
    
    def _auto_adjust_columns(self, sheet):
        """Auto-adjust column widths based on content."""
        for col_num in range(1, sheet.max_column + 1):
            max_length = 0
            column_letter = self._get_column_letter(col_num)
            
            for row_num in range(1, sheet.max_row + 1):
                cell = sheet.cell(row=row_num, column=col_num)
                try:
                    # Skip merged cells
                    if hasattr(cell, 'coordinate') and cell.coordinate in sheet.merged_cells:
                        continue
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            # Set width with some padding, but cap at reasonable maximum
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column_letter].width = adjusted_width
    
    def _create_header_style(self) -> Dict[str, Any]:
        """Create header cell styling."""
        header_config = self.output_config.get('header_style', {})
        
        return {
            'font': Font(
                bold=header_config.get('bold', True),
                color="FFFFFF"
            ),
            'fill': PatternFill(
                start_color=header_config.get('background_color', 'E6E6FA'),
                end_color=header_config.get('background_color', 'E6E6FA'),
                fill_type="solid"
            ),
            'alignment': Alignment(
                horizontal="center",
                vertical="center"
            )
        }
    
    def _create_hyperlink_style(self) -> Dict[str, Any]:
        """Create hyperlink cell styling."""
        hyperlink_color = self.output_config.get('hyperlink_color', '0000FF')
        
        return {
            'font': Font(
                color=hyperlink_color,
                underline="single"
            )
        }
    
    def _create_border_style(self) -> Border:
        """Create border styling."""
        thin_border = Side(border_style="thin", color="000000")
        return Border(
            left=thin_border,
            right=thin_border,
            top=thin_border,
            bottom=thin_border
        )
    
    def add_validation_summary_sheet(self, workbook_path: str, validation_reports: Dict[str, Any]):
        """Add a validation summary sheet to an existing workbook."""
        workbook = openpyxl.load_workbook(workbook_path)
        
        # Create validation summary sheet
        summary_sheet = workbook.create_sheet("Validation Summary")
        
        # Title
        summary_sheet['A1'] = "Validation Summary Report"
        summary_sheet['A1'].font = Font(size=16, bold=True)
        summary_sheet.merge_cells('A1:E1')
        
        current_row = 3
        
        # Overall statistics
        total_errors = sum(report.get_summary()['error_count'] for report in validation_reports.values())
        total_warnings = sum(report.get_summary()['warning_count'] for report in validation_reports.values())
        
        summary_sheet[f'A{current_row}'] = f"Total Errors: {total_errors}"
        summary_sheet[f'A{current_row}'].font = Font(color="FF0000" if total_errors > 0 else "008000")
        current_row += 1
        
        summary_sheet[f'A{current_row}'] = f"Total Warnings: {total_warnings}"
        summary_sheet[f'A{current_row}'].font = Font(color="FF8C00" if total_warnings > 0 else "008000")
        current_row += 2
        
        # Detailed reports for each supplier
        for supplier_name, report in validation_reports.items():
            summary_sheet[f'A{current_row}'] = f"=== {supplier_name} ==="
            summary_sheet[f'A{current_row}'].font = Font(bold=True)
            current_row += 1
            
            # Add validation report text
            report_text = self._create_validation_summary_text(report)
            for line in report_text.split('\n'):
                if line.strip():
                    summary_sheet[f'A{current_row}'] = line
                    current_row += 1
            
            current_row += 1
        
        # Auto-adjust columns
        self._auto_adjust_columns(summary_sheet)
        
        # Save workbook
        workbook.save(workbook_path)
        logging.info("Added validation summary sheet to workbook")
    
    def _create_validation_summary_text(self, validation_result) -> str:
        """Create a summary text from validation result."""
        lines = []
        summary = validation_result.get_summary()
        
        lines.append(f"Errors: {summary['error_count']}")
        lines.append(f"Warnings: {summary['warning_count']}")
        lines.append(f"Flagged Rows: {summary['flagged_rows']}")
        
        if validation_result.errors:
            lines.append("\nErrors:")
            for error in validation_result.errors[:5]:  # Limit to first 5
                lines.append(f"  - {error['message']}")
            if len(validation_result.errors) > 5:
                lines.append(f"  ... and {len(validation_result.errors) - 5} more")
        
        if validation_result.warnings:
            lines.append("\nWarnings:")
            for warning in validation_result.warnings[:5]:  # Limit to first 5
                lines.append(f"  - {warning['message']}")
            if len(validation_result.warnings) > 5:
                lines.append(f"  ... and {len(validation_result.warnings) - 5} more")
        
        return '\n'.join(lines)


def create_excel_generator(config: Dict[str, Any]) -> ExcelGenerator:
    """
    Create an Excel generator instance.
    
    Args:
        config: Configuration dictionary
    
    Returns:
        ExcelGenerator instance
    """
    return ExcelGenerator(config)


def generate_master_excel(supplier_data: Dict[str, pd.DataFrame], output_path: str, 
                         config: Dict[str, Any], validation_reports: Dict[str, Any] = None) -> str:
    """
    Generate a master Excel workbook with all supplier data.
    
    Args:
        supplier_data: Dictionary of supplier_name -> DataFrame
        output_path: Path for the output Excel file
        config: Configuration dictionary
        validation_reports: Optional validation reports for each supplier
    
    Returns:
        Path to the generated Excel file
    """
    generator = ExcelGenerator(config)
    return generator.generate_master_workbook(supplier_data, output_path, validation_reports)
